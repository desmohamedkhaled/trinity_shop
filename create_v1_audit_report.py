from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = 'Trinity_V1_Final_Audit_Report.docx'
BLUE = '083B68'
RED = 'B42318'
AMBER = '9A6700'
GREEN = '027A48'
GRAY = '475467'

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')

def add_text(cell, text, bold=False, color=None, size=9):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold; r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.7); sec.right_margin = Inches(.7)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10)
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
for name, size in [('Title', 25), ('Heading 1', 15), ('Heading 2', 11)]:
    st = styles[name]; st.font.name = 'Aptos Display'; st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0,0,0)
    st._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos Display')
    st._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos Display')

title = doc.add_paragraph(style='Title')
title.add_run('Trinity V1 Final Audit Report')
sub = doc.add_paragraph('Release readiness assessment for Trinity Christian Gift Shop')
sub.runs[0].font.color.rgb = RGBColor.from_string(GRAY); sub.runs[0].font.size = Pt(11)
sub.paragraph_format.space_after = Pt(12)

p = doc.add_paragraph()
p.add_run('Conclusion. ').bold = True
p.add_run('Trinity is not ready to launch. The application code passes its production build, type-check, and lint gates, but the live Supabase project is missing required checkout and Gift Finder schema, and it has no real catalog, occasions, or Gift Finder questions. Current V1 readiness is 55 percent.')

doc.add_heading('Audit scope and evidence', level=1)
doc.add_paragraph('The audit covered the storefront, shop and product paths, cart and checkout, wishlist, Gift Finder, WhatsApp, admin areas, authentication and authorization, Supabase, RLS, Storage, responsive behavior, SEO basics, and loading, error, and empty states. Supabase connectivity, public catalog data, settings, storage access, and the applicable database schema were checked against the configured live project. No claim of full live verification is made because the required live schema is absent.')

doc.add_heading('Readiness summary', level=1)
t = doc.add_table(rows=1, cols=3); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'
for cell, text in zip(t.rows[0].cells, ['Classification', 'Count', 'Assessment']):
    set_cell_shading(cell, BLUE); set_cell_margins(cell); add_text(cell, text, True, 'FFFFFF', 9)
rows = [
    ('Must fix before V1', '4', 'Live database and catalog blockers'),
    ('Needs test before V1', '5', 'Authenticated and end to end validation'),
    ('Ready', 'Core code paths', 'Build and code audit passed'),
    ('V1.1', '3', 'Non-blocking quality improvements'),
]
colors = [RED, AMBER, GREEN, GRAY]
for (label, count, assessment), color in zip(rows, colors):
    cells = t.add_row().cells
    for cell in cells: set_cell_margins(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    add_text(cells[0], label, True, color); add_text(cells[1], count, True); add_text(cells[2], assessment)

doc.add_heading('Must fix before V1', level=1)
for item in [
    'Apply the V1 schema repair migration. Live Supabase is missing the Gift Finder request fields and recommendation table, plus the create_checkout_order RPC. Checkout and customer requests cannot work until this is applied.',
    'Add and publish real products with stock. The one verified test product was removed during this audit, leaving the live catalog at zero products.',
    'Add and publish real occasions. Live count: zero.',
    'Add active Gift Finder questions and options. Live count: zero.',
]: add_bullet(doc, item)

doc.add_heading('Fixes completed during this audit', level=1)
for item in [
    'Corrected the storefront Gift Finder to use the public catalog endpoint instead of the admin-only endpoint.',
    'Allowed content managers to use the matching admin media endpoint.',
    'Changed shop categories to derive from the live product catalog rather than a hardcoded list.',
    'Removed fabricated Journal articles and replaced them with an honest empty state.',
    'Removed the verified live test product Test Trinity Cross Necklace from Supabase.',
    'Added an idempotent schema repair migration for the partial production deployment.',
]: add_bullet(doc, item)

doc.add_heading('Needs test before V1', level=1)
for item in [
    'Admin CRUD and role behavior with a real active administrator and a non-admin account.',
    'A full checkout with a real stocked product, including order creation, inventory decrement, admin visibility, and WhatsApp handoff.',
    'Admin media upload and delete against the Trinity Storage bucket.',
    'RLS behavior with anonymous, authenticated non-admin, and administrator sessions.',
    'Final mobile and desktop regression with the real catalog content.',
]: add_bullet(doc, item)

doc.add_heading('Ready', level=1)
for item in [
    'npm run build completed successfully.',
    'npx tsc --noEmit completed successfully.',
    'npm run lint completed with zero errors and three non-blocking warnings.',
    'Storefront routing, search, cart, wishlist, empty states, responsive navigation, and public catalog routes were audited.',
    'Server-side admin API authorization and the configured Supabase Storage bucket were inspected.',
    'Live settings confirm that WhatsApp checkout is enabled and a business number is configured.',
]: add_bullet(doc, item)

doc.add_heading('V1.1 improvements', level=1)
for item in [
    'Replace the two remaining img elements with optimized next/image usage.',
    'Resolve the anonymous default export lint warning in PostCSS configuration.',
    'Add automated end to end tests and request rate limiting.',
]: add_bullet(doc, item)

doc.add_heading('Launch checklist', level=1)
steps = [
    'Apply supabase/migrations/20260911_v1_schema_repair.sql to the live project.',
    'Verify that the checkout RPC and Gift Finder request fields are present.',
    'Create real products, occasions, and Gift Finder questions through Admin and publish them.',
    'Confirm an active admin_users record exists for the launch administrator.',
    'Place one real test order and verify its record, inventory update, admin visibility, and WhatsApp handoff.',
    'Deploy only after the tests above pass against the production Supabase project.',
]
for i, step in enumerate(steps, 1):
    p = doc.add_paragraph(style='List Number'); p.paragraph_format.space_after = Pt(4); p.add_run(step)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run('Trinity V1 Final Audit Report')
fr.font.size = Pt(8); fr.font.color.rgb = RGBColor.from_string(GRAY)

doc.core_properties.title = 'Trinity V1 Final Audit Report'
doc.core_properties.author = 'Trinity'
doc.save(OUT)
